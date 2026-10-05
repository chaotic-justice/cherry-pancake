import re
from dataclasses import dataclass
from io import BytesIO
from typing import BinaryIO, Sequence

import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from app.library.costco_store_catalog import (
    CostcoStoreCatalog,
    StoreCatalogError,
    StoreCatalogUpdate,
)
from app.library.costco_upload_policy import COSTCO_UPLOAD_POLICY
from app.library.utils import extract_mm_dd, extract_payment_id, get_today_date

EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATE_RANGE_ORDER_PATTERN = (
    r"\s*\d{1,2}/\d{1,2}/\d{2,4}\s*[-–—]\s*"
    r"\d{1,2}/\d{1,2}/\d{2,4}\s*"
)


@dataclass(frozen=True)
class InputFile:
    name: str
    body: BinaryIO


@dataclass(frozen=True)
class GeneratedWorkbook:
    body: bytes
    download_name: str
    media_type: str = EXCEL_MEDIA_TYPE


class CostcoInputError(ValueError):
    """A validation or document error safe to show to the user."""


@dataclass(frozen=True)
class _ReportMetadata:
    report_date: str | None
    payment_id: str | None


@dataclass(frozen=True)
class _ReportIdentity:
    sheet_title: str
    report_date: str
    check_number: str


@dataclass(frozen=True)
class _AnalyzedReport:
    source_name: str
    metadata: _ReportMetadata
    detail: pd.DataFrame
    summary: pd.DataFrame


def analyze_costco_reports(
    reports: Sequence[InputFile],
    store_update: InputFile | None = None,
) -> GeneratedWorkbook:
    """Validate Costco inputs and generate one analysis workbook in memory."""
    if not reports:
        raise CostcoInputError("Please upload at least one Costco PDF file.")
    if len(reports) > COSTCO_UPLOAD_POLICY.max_report_files:
        raise CostcoInputError(COSTCO_UPLOAD_POLICY.too_many_reports_message())

    store_catalog = _load_store_catalog(store_update)
    analyzed_reports = [
        _analyze_report(report, store_catalog) for report in reports
    ]
    return _render_workbook(analyzed_reports)


def _load_store_catalog(store_update: InputFile | None) -> CostcoStoreCatalog:
    update = None
    if store_update is not None and store_update.name:
        content = _read_limited(store_update, "Store mapping file")
        update = StoreCatalogUpdate(store_update.name, content)

    try:
        return CostcoStoreCatalog.load(update)
    except StoreCatalogError as error:
        raise CostcoInputError(str(error)) from error


def _analyze_report(
    report: InputFile,
    store_catalog: CostcoStoreCatalog,
) -> _AnalyzedReport:
    if not report.name or not report.name.lower().endswith(
        COSTCO_UPLOAD_POLICY.report_suffixes
    ):
        raise CostcoInputError("Costco report files must be PDFs.")

    content = _read_limited(report, report.name)
    try:
        with pdfplumber.open(BytesIO(content)) as pdf:
            if not pdf.pages:
                raise CostcoInputError(
                    f"No payment rows were found in {report.name}."
                )

            metadata = _extract_metadata(pdf.pages[0].extract_text() or "")
            rows: list[dict[str, str | float]] = []
            for page in pdf.pages:
                for table in page.extract_tables():
                    if not table or len(table) < 2:
                        continue
                    for row in table[1:]:
                        parsed = _parse_payment_row(row, report.name)
                        if parsed is not None:
                            rows.append(parsed)
    except CostcoInputError:
        raise
    except Exception as error:
        raise CostcoInputError(
            f"Could not read {report.name}. Check that it is a valid Costco PDF."
        ) from error

    if not rows:
        raise CostcoInputError(f"No payment rows were found in {report.name}.")

    detail = pd.DataFrame(rows)
    _assign_store_names(detail, store_catalog)
    summary = detail[["storeName", "amount"]].copy()
    summary = summary.groupby("storeName", as_index=False).sum()
    return _AnalyzedReport(report.name, metadata, detail, summary)


def _extract_metadata(text: str) -> _ReportMetadata:
    report_date = None
    payment_id = None
    for line in text.splitlines():
        if report_date is None and line.startswith("Date"):
            matched, value = extract_mm_dd(line)
            if matched:
                report_date = value
        elif payment_id is None and line.startswith("Payment"):
            matched, value = extract_payment_id(line)
            if matched:
                payment_id = value
    return _ReportMetadata(report_date, payment_id)


def _parse_payment_row(
    row: list[str | None],
    source_name: str,
) -> dict[str, str | float] | None:
    if not row or len(row) < 7:
        return None
    try:
        invoice = row[0] or ""
        order_number = row[1] or ""
        description = row[2] or ""
        payment_date = row[3] or ""
        amount_text = (row[6] or "0").replace(",", "").strip()
        if not amount_text or not invoice:
            return None
        return {
            "invoiceNumber": invoice.strip(),
            "orderNumber": order_number.strip(),
            "description": description.strip(),
            "date": payment_date.strip(),
            "amount": float(amount_text),
        }
    except (AttributeError, ValueError, TypeError, IndexError) as error:
        raise CostcoInputError(
            f"Could not parse a payment row in {source_name}."
        ) from error


def _assign_store_names(
    detail: pd.DataFrame,
    store_catalog: CostcoStoreCatalog,
) -> None:
    resolutions = detail["invoiceNumber"].map(store_catalog.resolve_invoice)
    detail["storeKey"] = resolutions.map(lambda resolution: resolution.store_key)
    detail["storeName"] = resolutions.map(lambda resolution: resolution.store_name)

    cosnext_mask = detail["amount"].lt(0) & detail["orderNumber"].str.fullmatch(
        DATE_RANGE_ORDER_PATTERN,
        na=False,
    )
    detail.loc[cosnext_mask, "storeName"] = "COSNEXT"


def _render_workbook(reports: Sequence[_AnalyzedReport]) -> GeneratedWorkbook:
    workbook = Workbook()
    for report in reports:
        identity = _report_identity(report)
        safe_title = re.sub(r"[\\*?:/\[\]]", "", identity.sheet_title)[:31]
        worksheet = workbook.create_sheet(title=safe_title)

        for row in dataframe_to_rows(report.detail, index=False, header=True):
            worksheet.append(row)
        worksheet.append([])
        for row in dataframe_to_rows(report.summary, index=False, header=True):
            worksheet.append(row)

        worksheet.append([])
        worksheet.append(["Total", report.summary["amount"].sum()])
        worksheet.append(["Date", identity.report_date])
        worksheet.append(["Check Number", identity.check_number])

    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])

    output = BytesIO()
    workbook.save(output)
    return GeneratedWorkbook(
        body=output.getvalue(),
        download_name=f"Costco_{get_today_date()}.xlsx",
    )


def _report_identity(report: _AnalyzedReport) -> _ReportIdentity:
    if report.metadata.report_date and report.metadata.payment_id:
        sheet_title = f"{report.metadata.report_date} #{report.metadata.payment_id}"
    else:
        sheet_title = report.source_name

    return _ReportIdentity(
        sheet_title=sheet_title,
        report_date=report.metadata.report_date or "Unknown",
        check_number=report.metadata.payment_id or "Unknown",
    )


def _read_limited(input_file: InputFile, label: str) -> bytes:
    content = input_file.body.read(COSTCO_UPLOAD_POLICY.max_file_bytes + 1)
    if len(content) > COSTCO_UPLOAD_POLICY.max_file_bytes:
        raise CostcoInputError(COSTCO_UPLOAD_POLICY.oversized_file_message(label))
    return content
