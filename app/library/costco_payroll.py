from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
import re
from typing import BinaryIO, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.library.costco_analysis import EXCEL_MEDIA_TYPE


MAX_FILES = 30
MAX_FILE_BYTES = 5 * 1024 * 1024
DATE_PATTERN = re.compile(r"(\d{1,2})[.-](\d{1,2})[.-](\d{2,4})")
OUTPUT_COLUMNS = [
    "Payroll Date",
    "Status",
    "Source File",
    "Source Sheet",
    "Source Row",
    "Show #",
    "Whs. #",
    "Location",
    "State",
    "Start Date",
    "End Date",
    "Sales Person",
    "# Days",
    "Advance Payment Amount",
    "Payment Date",
    "Net Sales",
    "Rate",
    "Commission Earned $",
    "Salary",
    "Reimbursement",
    "Commission Balance Due",
    "Prior Show Balance",
    "Costco Next",
    "Hotel",
    "Other 1",
    "Other 2",
    "Advance Payment",
    "Amount Paid",
    "Arrears",
]
HEADER_ALIASES = {
    "show": "Show #",
    "shownumber": "Show #",
    "whs": "Whs. #",
    "whsnumber": "Whs. #",
    "warehouse": "Whs. #",
    "warehousenumber": "Whs. #",
    "location": "Location",
    "state": "State",
    "startdate": "Start Date",
    "enddate": "End Date",
    "salesperson": "Sales Person",
    "days": "# Days",
    "ofdays": "# Days",
    "numberofdays": "# Days",
    "advancepaymentamount": "Advance Payment Amount",
    "date": "Payment Date",
    "paymentdate": "Payment Date",
    "datepaid": "Payment Date",
    "netsales": "Net Sales",
    "rate": "Rate",
    "commissionearned": "Commission Earned $",
    "salary": "Salary",
    "reimbursement": "Reimbursement",
    "reimbursements": "Reimbursement",
    "reimburesments": "Reimbursement",
    "commissionadded": "Commission Balance Due",
    "commissionbalancedue": "Commission Balance Due",
    "priorshowbalance": "Prior Show Balance",
    "costconext": "Costco Next",
    "costconextcommission": "Costco Next",
    "hotel": "Hotel",
    "other": "Other",
    "advancepayment": "Advance Payment",
    "amountdue": "Amount Paid",
    "amountpaid": "Amount Paid",
    "amounpaid": "Amount Paid",
    "arrears": "Arrears",
}
CURRENCY_COLUMNS = {
    "Advance Payment Amount",
    "Net Sales",
    "Commission Earned $",
    "Salary",
    "Reimbursement",
    "Commission Balance Due",
    "Prior Show Balance",
    "Costco Next",
    "Hotel",
    "Other 1",
    "Other 2",
    "Advance Payment",
    "Amount Paid",
    "Arrears",
}
DATE_COLUMNS = {"Payroll Date", "Start Date", "End Date", "Payment Date"}


@dataclass(frozen=True)
class PayrollFile:
    name: str
    body: BinaryIO


@dataclass(frozen=True)
class PayrollReport:
    body: bytes
    download_name: str
    master: tuple[dict, ...]
    review: tuple[dict, ...]
    source_count: int
    media_type: str = EXCEL_MEDIA_TYPE


class PayrollInputError(ValueError):
    """An input error safe to show to the user."""


def build_costco_payroll_master(files: Sequence[PayrollFile]) -> PayrollReport:
    if not files:
        raise PayrollInputError("Choose at least one Costco payroll XLSX file.")

    records: list[dict] = []
    for uploaded in files:
        records.extend(_read_workbook(uploaded))
    if not records:
        raise PayrollInputError("No Costco payroll rows were found in the uploaded files.")

    master, review = _resolve_records(records)
    return _write_workbook(master, review, records)


def _read_workbook(uploaded: PayrollFile) -> list[dict]:
    if not uploaded.name.lower().endswith(".xlsx"):
        raise PayrollInputError(f"{uploaded.name} must be an XLSX workbook.")
    source_date = _filename_date(uploaded.name)
    content = uploaded.body.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise PayrollInputError(f"{uploaded.name} must be 5 MB or smaller.")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    except Exception as error:
        raise PayrollInputError(
            f"Could not read {uploaded.name}. Check that it is a valid XLSX workbook."
        ) from error

    sheet = workbook.active
    header_rows = [
        row
        for row in range(1, min(sheet.max_row, 30) + 1)
        if _normalized(sheet.cell(row, 1).value) == "show"
        and "salesperson"
        in {_normalized(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)}
    ]
    if not header_rows:
        raise PayrollInputError(f"Could not find a payroll table in {uploaded.name}.")

    records = []
    for index, header_row in enumerate(header_rows):
        next_header = header_rows[index + 1] if index + 1 < len(header_rows) else min(sheet.max_row + 1, header_row + 50)
        columns = _header_columns(sheet, header_row)
        section = "Provisional" if "advance" in str(sheet.cell(header_row - 1, 1).value or "").lower() else "Final"
        for row_number in range(header_row + 1, next_header):
            values = {name: _typed_value(sheet.cell(row_number, column).value, name) for column, name in columns.items()}
            if not values.get("Sales Person") or not any(
                values.get(name) for name in ("Show #", "Whs. #", "Location")
            ):
                continue
            record = {name: values.get(name) for name in OUTPUT_COLUMNS[5:]}
            record.update(
                {
                    "Payroll Date": source_date,
                    "Status": section,
                    "Source File": uploaded.name,
                    "Source Sheet": sheet.title,
                    "Source Row": row_number,
                    "Decision": "",
                    "Review Reason": "",
                }
            )
            records.append(record)
    if not records:
        raise PayrollInputError(f"Could not find payroll rows in {uploaded.name}.")
    return records


def _header_columns(sheet, header_row: int) -> dict[int, str]:
    columns = {}
    other_count = 0
    for column in range(1, sheet.max_column + 1):
        canonical = HEADER_ALIASES.get(_normalized(sheet.cell(header_row, column).value))
        if canonical == "Other":
            other_count += 1
            canonical = f"Other {other_count}"
        if canonical:
            columns[column] = canonical
    return columns


def _resolve_records(records: list[dict]) -> tuple[list[dict], list[dict]]:
    groups: dict[tuple, list[dict]] = defaultdict(list)
    review = []
    for record in records:
        key = _record_key(record)
        if key is None:
            record["Decision"] = "Needs review"
            record["Review Reason"] = "Missing fields needed to identify the show."
            review.append(record)
        else:
            groups[key].append(record)

    master = []
    for group in groups.values():
        winning_rank = max(_rank(record) for record in group)
        candidates = [record for record in group if _rank(record) == winning_rank]
        signatures = {_signature(record) for record in candidates}
        if len(signatures) > 1:
            reason = "Conflicting records have the same payroll date and status."
            for record in group:
                record["Decision"] = "Needs review"
                record["Review Reason"] = reason
            review.extend(candidates)
            continue

        winner = min(candidates, key=lambda record: record["Source File"].lower().startswith("copy of "))
        winner["Decision"] = "Kept"
        winner["Status"] = "Final" if winner["Status"] == "Final" else "Provisional only"
        master.append(winner)
        for record in group:
            if record is winner:
                continue
            if record in candidates:
                record["Decision"] = "Exact duplicate"
            else:
                record["Decision"] = f"Superseded by {winner['Source File']}"

    master.sort(key=lambda record: (record["Payroll Date"], _text(record.get("Sales Person")), _text(record.get("Show #"))))
    review.sort(key=lambda record: (record["Payroll Date"], record["Source File"], record["Source Row"]))
    records.sort(key=lambda record: (record["Payroll Date"], record["Source File"], record["Source Row"]))
    return master, review


def _record_key(record: dict) -> tuple | None:
    warehouse = _key_value(record.get("Whs. #"))
    salesperson = _key_value(record.get("Sales Person"))
    start_date = _key_value(record.get("Start Date"))
    if warehouse and salesperson and start_date:
        return "warehouse", warehouse, salesperson, start_date

    show = _key_value(record.get("Show #"))
    location = _key_value(record.get("Location"))
    end_date = _key_value(record.get("End Date"))
    if show and location and salesperson and (start_date or end_date):
        return "show", show, location, salesperson, start_date, end_date
    return None


def _rank(record: dict) -> tuple[int, date, int]:
    # ponytail: Drive "Copy of" files are treated as later revisions; add
    # explicit version metadata if filenames stop reflecting revision order.
    copy_revision = int(record["Source File"].lower().startswith("copy of "))
    return (1 if record["Status"] == "Final" else 0, record["Payroll Date"], copy_revision)


def _signature(record: dict) -> tuple:
    return tuple(_key_value(record.get(column)) for column in OUTPUT_COLUMNS[5:])


def _write_workbook(master: list[dict], review: list[dict], records: list[dict]) -> PayrollReport:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_sheet(workbook, "Payroll Master", master, OUTPUT_COLUMNS)
    _append_sheet(
        workbook,
        "Needs Review",
        review,
        ["Review Reason", *OUTPUT_COLUMNS],
    )
    _append_sheet(
        workbook,
        "Source Rows",
        records,
        ["Decision", *OUTPUT_COLUMNS],
    )
    output = BytesIO()
    workbook.save(output)
    latest_date = max(record["Payroll Date"] for record in records)
    return PayrollReport(
        body=output.getvalue(),
        download_name=f"Costco_Payroll_Master_{latest_date.isoformat()}.xlsx",
        master=tuple(master),
        review=tuple(review),
        source_count=len(records),
    )


def _append_sheet(workbook: Workbook, title: str, records: list[dict], columns: list[str]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for record in records:
        sheet.append([record.get(column) for column in columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, name in enumerate(columns, start=1):
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            if name in DATE_COLUMNS:
                cell.number_format = "mm/dd/yy"
            elif name in CURRENCY_COLUMNS:
                cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
            elif name == "Rate":
                cell.number_format = "0.0%"
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(max(len(value) for value in values) + 2, 28)


def _filename_date(filename: str) -> date:
    match = DATE_PATTERN.search(filename)
    if not match:
        raise PayrollInputError(f"{filename} must include a payroll date such as 10.04.26.")
    month, day, year = (int(value) for value in match.groups())
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError as error:
        raise PayrollInputError(f"{filename} contains an invalid payroll date.") from error


def _typed_value(value: object, column: str) -> object:
    if column in DATE_COLUMNS:
        return _parse_date(value)
    if isinstance(value, str):
        return value.strip()
    return value


def _parse_date(value: object) -> object:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str) or not value.strip():
        return value
    match = DATE_PATTERN.search(value.strip())
    if not match:
        return value.strip()
    month, day, year = (int(part) for part in match.groups())
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return value.strip()


def _normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _key_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value, 8)
    if isinstance(value, str):
        return " ".join(value.split()).casefold()
    return value


def _text(value: object) -> str:
    return str(_key_value(value))
