import unittest
from io import BytesIO

from openpyxl import load_workbook

from app.library.costco_analysis import (
    CostcoInputError,
    InputFile,
    analyze_costco_reports,
)
from tests.costco_pdf_fixture import make_costco_pdf


def _footer_value(sheet, label: str):
    return next(
        row[1].value
        for row in sheet.iter_rows()
        if row[0].value == label
    )


class CostcoAnalysisTest(unittest.TestCase):
    def test_generates_workbook_from_real_pdf_bytes(self):
        pdf = make_costco_pdf(
            [
                [
                    "0141441151",
                    "8/3/26 - 8/30/26",
                    "",
                    "09/22/2026",
                    "-5,423.93",
                    "0.00",
                    "-5,423.93",
                ],
                [
                    "0141441152",
                    "8/3/26 - 8/30/26",
                    "",
                    "09/22/2026",
                    "100.00",
                    "0.00",
                    "100.00",
                ],
            ],
            metadata_lines=("Payment: 456", "Date: 09/24/2026"),
        )

        generated = analyze_costco_reports([InputFile("payment.pdf", BytesIO(pdf))])

        workbook = load_workbook(BytesIO(generated.body))
        sheet = workbook[workbook.sheetnames[0]]
        self.assertEqual(sheet["F2"].value, "0141")
        self.assertEqual(sheet["G2"].value, "COSNEXT")
        self.assertEqual(sheet["G3"].value, "C7#0141")
        self.assertEqual(_footer_value(sheet, "Date"), "09-24")
        self.assertEqual(_footer_value(sheet, "Check Number"), "456")

    def test_supports_duplicate_payment_ids_and_store_updates(self):
        pdf = make_costco_pdf(
            [["0999441151", "1", "Item", "09/22/2026", "10", "0", "10"]]
        )
        reports = [
            InputFile("first.pdf", BytesIO(pdf)),
            InputFile("second.pdf", BytesIO(pdf)),
        ]
        store_update = InputFile(
            "stores.csv",
            BytesIO(b"C2#999,COSTCO,#0999\n"),
        )

        generated = analyze_costco_reports(reports, store_update)

        workbook = load_workbook(BytesIO(generated.body))
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertEqual(workbook[workbook.sheetnames[0]]["G2"].value, "C2#0999")
        self.assertEqual(workbook[workbook.sheetnames[1]]["G2"].value, "C2#0999")

    def test_preserves_partial_and_missing_metadata_explicitly(self):
        rows = [["0141441151", "1", "Item", "09/22/2026", "10", "0", "10"]]
        cases = [
            (
                "date-only.pdf",
                ("Date: 09/24/2026",),
                "09-24",
                "Unknown",
            ),
            (
                "payment-only.pdf",
                ("Payment: 456",),
                "Unknown",
                "456",
            ),
            (
                "missing-metadata.pdf",
                (),
                "Unknown",
                "Unknown",
            ),
        ]

        for filename, metadata_lines, expected_date, expected_check in cases:
            with self.subTest(filename=filename):
                pdf = make_costco_pdf(rows, metadata_lines=metadata_lines)
                generated = analyze_costco_reports(
                    [InputFile(filename, BytesIO(pdf))]
                )
                workbook = load_workbook(BytesIO(generated.body))
                sheet = workbook[filename]

                self.assertEqual(_footer_value(sheet, "Date"), expected_date)
                self.assertEqual(
                    _footer_value(sheet, "Check Number"),
                    expected_check,
                )

    def test_first_duplicate_metadata_values_win(self):
        pdf = make_costco_pdf(
            [["0141441151", "1", "Item", "09/22/2026", "10", "0", "10"]],
            metadata_lines=(
                "Payment: 456",
                "Payment: 999",
                "Date: 09/24/2026",
                "Date: 10/31/2026",
            ),
        )

        generated = analyze_costco_reports(
            [InputFile("duplicates.pdf", BytesIO(pdf))]
        )
        workbook = load_workbook(BytesIO(generated.body))
        sheet = workbook["09-24 #456"]

        self.assertEqual(_footer_value(sheet, "Date"), "09-24")
        self.assertEqual(_footer_value(sheet, "Check Number"), "456")

    def test_rejects_invalid_inputs_at_the_boundary(self):
        with self.assertRaisesRegex(CostcoInputError, "at least one"):
            analyze_costco_reports([])

        with self.assertRaisesRegex(CostcoInputError, "must be PDFs"):
            analyze_costco_reports([InputFile("report.txt", BytesIO(b"text"))])

        pdf = make_costco_pdf(
            [["0141441151", "1", "Item", "09/22/2026", "10", "0", "10"]]
        )
        with self.assertRaisesRegex(CostcoInputError, "CSV, XLS, or XLSX"):
            analyze_costco_reports(
                [InputFile("report.pdf", BytesIO(pdf))],
                InputFile("stores.txt", BytesIO(b"not-a-store-update")),
            )

    def test_enforces_the_shared_upload_count_and_size_limits(self):
        reports = [
            InputFile(f"report-{index}.pdf", BytesIO())
            for index in range(31)
        ]
        with self.assertRaisesRegex(CostcoInputError, "no more than 30"):
            analyze_costco_reports(reports)

        oversized = b"x" * (5 * 1024 * 1024 + 1)
        with self.assertRaisesRegex(CostcoInputError, "5 MB or smaller"):
            analyze_costco_reports(
                [InputFile("oversized.pdf", BytesIO(oversized))]
            )

        with self.assertRaisesRegex(CostcoInputError, "5 MB or smaller"):
            analyze_costco_reports(
                [InputFile("report.pdf", BytesIO(b"unused"))],
                InputFile("stores.csv", BytesIO(oversized)),
            )


if __name__ == "__main__":
    unittest.main()
