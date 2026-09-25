from io import BytesIO
import unittest

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.library.reconciliation import MonthFrames, reconcile_months
from app.main import app
from app.routes.raccoon import build_reconciliation_files


def workbook(month: int, check: int) -> bytes:
    output = BytesIO()
    date = pd.Timestamp(2026, month, 3)
    vnb = pd.DataFrame(
        [["40395782", date, check, "CHECK", 100.0, None]],
        columns=["Account Number", "Post Date", "Check", "Description", "Debit", "Credit"],
    )
    ap = pd.DataFrame(
        [[check, date, "VENDOR", "Example Vendor", 100.0]],
        columns=["Check Number", "Check Date", "Vendor Number", "Name", "Check Amount"],
    )
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        vnb.to_excel(writer, sheet_name="VNB", index=False)
        ap.to_excel(writer, sheet_name="AP", index=False, startrow=1)
    return output.getvalue()


def ar_matching_workbook() -> tuple[bytes, bytes]:
    recon_output = BytesIO()
    vnb = pd.DataFrame(
        [
            ["1", pd.Timestamp(2026, 1, 5), "", "PAYMENT REF12345", None, 100.0],
            ["1", pd.Timestamp(2026, 1, 6), "", "DEPOSIT", None, 300.0],
            ["1", pd.Timestamp(2026, 1, 9), "", "ACH CREDIT CCD CMPY ID: 4270465600 TRANSFER", None, 400.0],
            ["1", pd.Timestamp(2026, 1, 10), "", "ACH CREDIT CCD CMPY ID: 4270465600 BADREF", None, 500.0],
        ],
        columns=["Account Number", "Post Date", "Check", "Description", "Debit", "Credit"],
    )
    ap = pd.DataFrame(
        [["1", pd.Timestamp(2026, 1, 1), "VENDOR", "Example", 1.0]],
        columns=["Check Number", "Check Date", "Vendor Number", "Name", "Check Amount"],
    )
    with pd.ExcelWriter(recon_output, engine="openpyxl") as writer:
        vnb.to_excel(writer, sheet_name="VNB", index=False)
        ap.to_excel(writer, sheet_name="AP", index=False)

    ar_output = BytesIO()
    check = pd.DataFrame(
        [
            [pd.Timestamp(1936, 2, 27), "Reference customer", "REF12345", None, "", None, "INV1", 100.0, None],
            [pd.Timestamp(2026, 1, 5), "Deposit customer", "", None, "D100", None, "INV2", 100.0, None],
            [pd.Timestamp(2026, 1, 5), "Deposit customer", "", None, "D100", None, "INV3", 200.0, None],
            [pd.Timestamp(2026, 1, 10), "Conflict customer", "BADREF", None, "", None, "INV4", 499.0, None],
        ]
    )
    internet = pd.DataFrame(
        [
            [pd.Timestamp(2026, 1, 5), "INTERNET CC", "", None, "", "I1", 400.0, None],
            [pd.Timestamp(2026, 1, 10), "INTERNET CC", "", None, "", "I2", 500.0, None],
        ]
    )
    with pd.ExcelWriter(ar_output, engine="openpyxl") as writer:
        check.to_excel(writer, sheet_name="CHECK", index=False, header=False)
        internet.to_excel(writer, sheet_name="INTERNET", index=False, header=False)
    return recon_output.getvalue(), ar_output.getvalue()


class RaccoonTest(unittest.TestCase):
    def test_infers_and_sorts_months_without_filename_conventions(self):
        files = build_reconciliation_files(
            [("anything.xlsx", workbook(2, 102)), ("untitled.xlsx", workbook(1, 101))]
        )

        self.assertEqual(
            [filename for filename, _ in files],
            ["raccoon-2026-01.xlsx", "raccoon-2026-02.xlsx"],
        )
        january = pd.read_excel(BytesIO(files[0][1]), sheet_name="Processed AP")
        self.assertEqual(january.loc[0, "Source Row"], 3)
        self.assertEqual(january.loc[0, "Match Status"], "matched")
        self.assertEqual(january.loc[0, "Match Rule"], "Unique amount")
        self.assertIn("only available VNB transaction", january.loc[0, "Match Details"])
        output = load_workbook(BytesIO(files[0][1]))
        ap_sheet = output["Processed AP"]
        self.assertEqual(ap_sheet.freeze_panes, "A2")
        self.assertEqual(ap_sheet.auto_filter.ref, ap_sheet.dimensions)
        self.assertTrue(
            {"AP Check Number", "AP Check Date", "Vendor Name", "AP Amount"}.issubset(
                cell.value for cell in ap_sheet[1]
            )
        )
        ap_status_column = next(
            cell.column for cell in ap_sheet[1] if cell.value == "Match Status"
        )
        self.assertTrue(ap_sheet.cell(2, ap_status_column).fill.fgColor.rgb.endswith("90EE90"))

    def test_ap_rules_use_check_number_then_date_for_repeated_amounts(self):
        dates = [pd.Timestamp(2026, 1, day).date() for day in (5, 6)]
        vnb = pd.DataFrame(
            [
                [2, dates[0], "1001", "CHECK", 100.0, None],
                [3, dates[1], "1002", "CHECK", 100.0, None],
            ],
            columns=["sourceRow", "postDate", "check", "description", "debit", "credit"],
        )
        ap = pd.DataFrame(
            [
                [3, "1001", dates[1], "A", "Vendor A", 100.0],
                [4, "9999", dates[1], "B", "Vendor B", 100.0],
            ],
            columns=["sourceRow", "checkNumber", "checkDate", "vendorNumber", "name", "checkAmount"],
        )

        result = reconcile_months([MonthFrames("2026-01", ap=ap, vnb=vnb)])[0].ap

        self.assertEqual(result["status"].tolist(), ["matched", "matched"])
        self.assertEqual(
            result["matchRule"].tolist(),
            ["Amount and check number", "Amount and date"],
        )
        self.assertTrue(result["notes"].str.contains("because", regex=False).all())

    def test_net_pr_split_reserves_vnb_rows_across_months(self):
        jan_date = pd.Timestamp(2026, 1, 31).date()
        feb_date = pd.Timestamp(2026, 2, 1).date()
        jan = MonthFrames(
            "2026-01",
            ap=pd.DataFrame(
                [
                    [3, "1", jan_date, "NET PR", "Payroll A", 100.0],
                    [4, "2", jan_date, "NET PR", "Payroll B", 100.0],
                ],
                columns=["sourceRow", "checkNumber", "checkDate", "vendorNumber", "name", "checkAmount"],
            ),
            vnb=pd.DataFrame(
                [[2, jan_date, "", "ADP WAGE", 80.0, None]],
                columns=["sourceRow", "postDate", "check", "description", "debit", "credit"],
            ),
        )
        feb = MonthFrames(
            "2026-02",
            ap=pd.DataFrame(
                [[3, "", feb_date, "OTHER", "Other payment", 20.0]],
                columns=["sourceRow", "checkNumber", "checkDate", "vendorNumber", "name", "checkAmount"],
            ),
            vnb=pd.DataFrame(
                [[2, feb_date, "", "OTHER", 20.0, None]],
                columns=["sourceRow", "postDate", "check", "description", "debit", "credit"],
            ),
        )

        january, february = reconcile_months([jan, feb])

        self.assertEqual(january.ap["status"].tolist(), ["split match", "no match"])
        self.assertEqual(february.ap.loc[0, "status"], "no match")
        self.assertEqual(february.vnb.loc[0, "status"], "partial match")

    def test_february_vendor_names_do_not_select_vnb_rows_by_position(self):
        bank_rows = [
            [index + 2, pd.Timestamp(2026, 2, 1).date(), "", "OTHER", None, None]
            for index in range(19)
        ]
        bank_rows[0][4] = 6000.0
        bank_rows[1][4] = 6000.0
        bank_rows[18][4] = 999.0
        month = MonthFrames(
            "2026-02",
            ap=pd.DataFrame(
                [[3, "DC2602", pd.Timestamp(2026, 2, 28).date(), "DANY", "DANY CHALHOUB", 6000.0]],
                columns=["sourceRow", "checkNumber", "checkDate", "vendorNumber", "name", "checkAmount"],
            ),
            vnb=pd.DataFrame(
                bank_rows,
                columns=["sourceRow", "postDate", "check", "description", "debit", "credit"],
            ),
        )

        result = reconcile_months([month])[0]

        self.assertEqual(result.ap.loc[0, "status"], "no match")
        self.assertEqual(result.ap.loc[0, "matchRule"], "Ambiguous amount")
        self.assertTrue(result.vnb["status"].eq("no match").all())

    def test_ar_rule_order_and_html_review(self):
        recon, ar = ar_matching_workbook()
        files = build_reconciliation_files(
            [("january.xlsx", recon)], [("01JANUARY2026.xlsx", ar)]
        )
        workbook_bytes = files[0][1]
        vnb = pd.read_excel(BytesIO(workbook_bytes), sheet_name="Processed VNB")
        output = load_workbook(BytesIO(workbook_bytes))

        self.assertEqual(
            vnb["AR Match Rule"].tolist(),
            [
                "Embedded reference and exact amount",
                "Batch/deposit total, source, amount, and date",
                "Source, exact amount, and 0–4 day lag",
                "Referenced amount differs",
            ],
        )
        self.assertIn("does not equal", vnb.loc[3, "AR Match Details"])
        self.assertEqual(vnb.loc[0, "AR Date"], "1936-02-27")
        self.assertIn("expected 0–4 days", vnb.loc[0, "Date Warning"])
        self.assertIn("AR Matches", output.sheetnames)
        ar_sheet = output["AR Matches"]
        ar_status_column = next(
            cell.column for cell in ar_sheet[1] if cell.value == "AR Match Status"
        )
        self.assertTrue(ar_sheet.cell(3, ar_status_column).fill.fgColor.rgb.endswith("93C5FD"))

        response = TestClient(app).post(
            "/raccoon",
            files=[
                ("files", ("january.xlsx", recon, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
                ("ar_files", ("01JANUARY2026.xlsx", ar, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ],
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Matching results", response.text)
        self.assertIn("Download 2026-01 Excel", response.text)
        self.assertIn('download="raccoon-2026-01.xlsx"', response.text)
        self.assertNotIn("application/zip", response.text)
        self.assertIn("Needs Review", response.text)
        self.assertIn("Date warnings", response.text)
        self.assertIn("Verify date", response.text)
        self.assertIn("1936-02-27", response.text)
        self.assertIn("1 input date issue", response.text)
        self.assertIn("outside the uploaded reconciliation range", response.text)


if __name__ == "__main__":
    unittest.main()
