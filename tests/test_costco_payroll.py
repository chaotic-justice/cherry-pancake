import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.library.costco_payroll import (
    PayrollFile,
    PayrollInputError,
    build_costco_payroll_master,
)


HEADERS = [
    "Show #",
    "Whs. #",
    "Location",
    "State",
    "Start Date",
    "End Date",
    "Sales Person",
    "# Days",
    "Advance Payment Amount",
    "Payment Date",
    "Net sales",
    "Rate",
    "Commission Earned $",
    "Salary",
    "Reimbursement",
    "Commission Balance Due",
    "Costco Next",
    "*Other",
    "*Other",
    "Amount Paid",
    "Arrears",
]


def workbook_bytes(final_rows=(), advance_rows=(), stale_row=None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A2"] = "Costco Payroll"
    sheet.append([])
    sheet.append(HEADERS)
    for row in final_rows:
        sheet.append(row)
    sheet["A18"] = "ADVANCE PAYMENTS"
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(19, column, header)
    for row in advance_rows:
        sheet.append(row)
    if stale_row:
        stale = workbook.create_sheet("paid")
        stale["A2"] = "Costco Payroll"
        for column, header in enumerate(HEADERS, start=1):
            stale.cell(3, column, header)
        for column, value in enumerate(stale_row, start=1):
            stale.cell(4, column, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def row(show, warehouse, person, start, net_sales=None, commission=None):
    values = [None] * len(HEADERS)
    values[0:8] = [show, warehouse, "Location", "NY", start, "10.11.26", person, 10]
    values[10:16] = [net_sales, 0.23, commission, 1500, 1500, max((commission or 0) - 3000, 0)]
    values[19] = 3000
    return values


class CostcoPayrollTest(unittest.TestCase):
    def test_later_final_row_supersedes_advance_and_ignores_stale_tabs(self):
        advance = row(100, 200, "Jane", "09.18.26")
        final = row(100, 200, "Jane", "09.18.26", 20_000, 4_600)
        next_advance = row(101, 201, "Alex", "10.02.26")
        files = [
            PayrollFile(
                "09.18.26.xlsx",
                BytesIO(workbook_bytes(advance_rows=[advance])),
            ),
            PayrollFile(
                "10.04.26.xlsx",
                BytesIO(
                    workbook_bytes(
                        final_rows=[final],
                        advance_rows=[next_advance],
                        stale_row=row(999, 999, "Stale", "01.01.26"),
                    )
                ),
            ),
        ]

        generated = build_costco_payroll_master(files)
        workbook = load_workbook(BytesIO(generated.body), data_only=True)
        master = workbook["Payroll Master"]
        source = workbook["Source Rows"]

        self.assertEqual(master.max_row, 3)
        headers = [cell.value for cell in master[1]]
        show_column = headers.index("Show #") + 1
        status_column = headers.index("Status") + 1
        self.assertEqual(
            {(master.cell(row_number, show_column).value, master.cell(row_number, status_column).value) for row_number in range(2, 4)},
            {(100, "Final"), (101, "Provisional only")},
        )
        decisions = [source.cell(row_number, 1).value for row_number in range(2, source.max_row + 1)]
        self.assertTrue(any(str(value).startswith("Superseded by 10.04.26") for value in decisions))
        self.assertNotIn(999, [master.cell(row_number, show_column).value for row_number in range(2, master.max_row + 1)])

    def test_copy_named_file_wins_as_later_same_date_revision(self):
        original = row(100, 200, "Jane", "01.20.26", 10_000, 2_300)
        revision = row(100, 200, "Jane", "01.20.26", 12_000, 2_760)
        generated = build_costco_payroll_master(
            [
                PayrollFile("01.26.26.xlsx", BytesIO(workbook_bytes(final_rows=[original]))),
                PayrollFile("Copy of 01.26.26.xlsx", BytesIO(workbook_bytes(final_rows=[revision]))),
            ]
        )
        workbook = load_workbook(BytesIO(generated.body), data_only=True)
        master = workbook["Payroll Master"]
        headers = [cell.value for cell in master[1]]

        self.assertEqual(master.cell(2, headers.index("Net Sales") + 1).value, 12_000)
        self.assertEqual(master.cell(2, headers.index("Source File") + 1).value, "Copy of 01.26.26.xlsx")

    def test_rejects_invalid_inputs(self):
        with self.assertRaisesRegex(PayrollInputError, "at least one"):
            build_costco_payroll_master([])
        with self.assertRaisesRegex(PayrollInputError, "must be an XLSX"):
            build_costco_payroll_master([PayrollFile("payroll.csv", BytesIO())])


if __name__ == "__main__":
    unittest.main()
