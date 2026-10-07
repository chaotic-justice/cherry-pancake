import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from app.library.sales_analysis import SalesInputError, build_sales_report


def workbook_bytes(rows: list[list[object]]) -> bytes:
    output = BytesIO()
    pd.DataFrame(rows, columns=list("abcde")).to_excel(output, index=False)
    return output.getvalue()


class SalesAnalysisTest(unittest.TestCase):
    def test_rejects_duplicate_summary_metrics(self):
        content = workbook_bytes(
            [
                ["Salesperson Jane", None, None, None, None],
                ["Period To Date:", 10, None, None, None],
                ["Period To Date:", 20, None, None, None],
                ["Period To Date:", 30, None, None, None],
                ["Period total", 60, None, None, None],
                ["Year total", 0, None, None, None],
                ["Prior total", 0, None, None, None],
                ["End", None, None, None, None],
            ]
        )

        with self.assertRaisesRegex(SalesInputError, "distinct"):
            build_sales_report(content)

    def test_keeps_nonzero_metrics_that_cancel(self):
        content = workbook_bytes(
            [
                ["Salesperson Jane", None, None, None, None],
                ["Period To Date:", 10, None, None, None],
                ["Year To Date:", -10, None, None, None],
                ["Prior Year:", 0, None, None, None],
                ["Period total", 10, None, None, None],
                ["Year total", -10, None, None, None],
                ["Prior total", 0, None, None, None],
                ["End", None, None, None, None],
            ]
        )

        sheet = load_workbook(BytesIO(build_sales_report(content))).active

        self.assertEqual(
            [sheet.cell(2, column).value for column in range(1, 5)],
            ["Jane", 10, -10, 0],
        )


if __name__ == "__main__":
    unittest.main()
