import unittest
from io import BytesIO

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.costco_pdf_fixture import make_costco_pdf


class EndpointTest(unittest.TestCase):
    client = TestClient(app)

    def test_costco_requires_a_pdf(self):
        response = self.client.post(
            "/costco",
            files={"files": ("report.txt", b"not a PDF", "text/plain")},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("must be PDFs", response.text)

    def test_costco_returns_generated_workbook(self):
        pdf = make_costco_pdf(
            [["000636ABCDEF", "1", "Item", "01/02", "10", "0", "10"]],
            metadata_lines=("Date: 01/02/2026", "Payment: 123"),
        )

        response = self.client.post(
            "/costco",
            files={"files": ("payment.pdf", pdf, "application/pdf")},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("Costco_", response.headers["content-disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook[workbook.sheetnames[0]]["G2"].value, "C4#0636")

    def test_sales_requires_an_excel_workbook(self):
        response = self.client.post(
            "/sales",
            files={"file": ("sales.txt", b"not a workbook", "text/plain")},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("must be an Excel workbook", response.text)

    def test_sales_rejects_an_unexpected_workbook_layout(self):
        workbook = BytesIO()
        pd.DataFrame([["only", 1, 2, 3, 4]], columns=list("abcde")).to_excel(
            workbook, index=False
        )

        response = self.client.post(
            "/sales",
            files={
                "file": (
                    "sales.xlsx",
                    workbook.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("summary rows", response.text)

    def test_sales_accepts_the_expected_workbook_layout(self):
        workbook = BytesIO()
        pd.DataFrame(
            [
                ["Salesperson Jane", None, None, None, None],
                ["Period To Date:", 10, None, None, None],
                ["Year To Date:", 20, None, None, None],
                ["Prior Year:", 30, None, None, None],
                ["Period total", 10, None, None, None],
                ["Year total", 20, None, None, None],
                ["Prior total", 30, None, None, None],
                ["End", None, None, None, None],
            ],
            columns=list("abcde"),
        ).to_excel(workbook, index=False)

        response = self.client.post(
            "/sales",
            files={
                "file": (
                    "sales.xlsx",
                    workbook.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            load_workbook(BytesIO(response.content)).sheetnames,
            ["Sales Report"],
        )

    def test_templates_load_static_assets(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertIn('/static/index.css', response.text)
        self.assertEqual(self.client.get("/static/base.css").status_code, 200)

    def test_home_page_renders_all_tool_rows(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.text.count('class="tool-row"'), 3)
        self.assertIn('<a href="/costco" class="tool-row">', response.text)
        self.assertIn('<a href="/sales" class="tool-row">', response.text)
        self.assertIn('<a href="/raccoon" class="tool-row">', response.text)

    def test_sales_page_renders_the_upload_workflow(self):
        response = self.client.get("/sales")

        self.assertEqual(response.status_code, 200)
        self.assertIn('<form class="workflow" action="/sales"', response.text)
        self.assertIn('<span class="file-picker-copy">', response.text)
        self.assertIn('id="analyze-btn" disabled', response.text)

    def test_raccoon_progressively_enhances_form_submission(self):
        page = self.client.get("/raccoon")
        script = self.client.get("/static/raccoon.js")

        self.assertIn('id="raccoon-form"', page.text)
        self.assertIn('id="request-status"', page.text)
        self.assertIn("event.preventDefault()", script.text)
        self.assertIn("new FormData(form)", script.text)
        self.assertIn('querySelector("main").replaceWith(nextMain)', script.text)

    def test_costco_page_explains_optional_store_update(self):
        response = self.client.get("/costco")

        self.assertEqual(response.status_code, 200)
        self.assertIn("built-in Costco store list", response.text)
        self.assertNotRegex(
            response.text,
            r'<input[^>]*id="store-input"[^>]*\brequired\b',
        )

    def test_costco_form_uses_the_server_upload_contract(self):
        response = self.client.get("/costco")
        script = self.client.get("/static/costco.js")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(script.status_code, 200)
        self.assertIn('data-max-report-files="30"', response.text)
        self.assertIn('data-max-file-bytes="5242880"', response.text)
        self.assertIn('data-max-file-size-mb="5"', response.text)
        self.assertIn('data-report-suffixes=".pdf"', response.text)
        self.assertIn('data-store-update-suffixes=".csv,.xls,.xlsx"', response.text)
        self.assertIn("max 30, 5 MB each", response.text)
        self.assertIn(".csv,.xls,.xlsx", response.text)
        self.assertNotIn("100 MB", response.text)
        self.assertIn("analysisForm.dataset.maxReportFiles", script.text)
        self.assertIn("analysisForm.dataset.maxFileBytes", script.text)
        self.assertIn("analysisForm.dataset.reportSuffixes", script.text)
        self.assertIn("analysisForm.dataset.storeUpdateSuffixes", script.text)

    def test_openapi_declares_real_response_types(self):
        paths = app.openapi()["paths"]

        self.assertIn("text/html", paths["/"]["get"]["responses"]["200"]["content"])
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            paths["/sales"]["post"]["responses"]["200"]["content"],
        )


if __name__ == "__main__":
    unittest.main()
